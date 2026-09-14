from pathlib import Path

from docx import Document


document_path = Path.home() / "Documents" / "BlockBug_Software_Requirements_Specification.docx"
document = Document(document_path)

content = "\n".join(paragraph.text for paragraph in document.paragraphs)
for table in document.tables:
    for row in table.rows:
        content += "\n" + " | ".join(cell.text for cell in row.cells)

required_text = [
    "Ethereum Sepolia",
    "Alchemy",
    "Vercel",
    "Render",
    "Identified Gap",
    "Functional Requirements",
    "Appendix C",
]

print(f"File: {document_path}")
print(f"Size: {document_path.stat().st_size} bytes")
print(f"Paragraphs: {len(document.paragraphs)}")
print(f"Tables: {len(document.tables)}")
print(f"Sections: {len(document.sections)}")

missing = [text for text in required_text if text not in content]
if missing:
    raise SystemExit(f"Missing required text: {', '.join(missing)}")

print("Required content: PASS")
print("DOCX integrity: PASS")
