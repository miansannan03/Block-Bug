from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = Path(r"C:\Users\ASUS\Documents\BlockBug_Software_Requirements_Specification.docx")
PROJECT_NAME = "BlockBug"
VERSION = "1.0"
AUTHOR = "Mian Sannan"
ORGANIZATION = "Academic Software Project"
DATE = "June 10, 2026"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_toc(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and select Update Field to generate the table of contents."
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_separate, placeholder, fld_char_end])


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.08
        paragraph.add_run(item)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_requirement_table(doc, requirements):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Requirement ID", "Requirement", "Priority", "Verification"]
    for index, heading in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = heading
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    set_repeat_table_header(table.rows[0])
    for req_id, requirement, priority, verification in requirements:
        cells = table.add_row().cells
        values = [req_id, requirement, priority, verification]
        for index, value in enumerate(values):
            cells[index].text = value
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    doc.add_paragraph()


def add_two_column_table(doc, rows, headers=("Item", "Description"), widths=None):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, heading in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = heading
        set_cell_shading(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    set_repeat_table_header(table.rows[0])
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    if widths:
        for row in table.rows:
            row.cells[0].width = widths[0]
            row.cells[1].width = widths[1]
    doc.add_paragraph()
    return table


def add_feature(doc, number, title, description, priority, sequence, requirements):
    add_heading(doc, f"{number} {title}", 2)
    add_heading(doc, f"{number}.1 Description and Priority", 3)
    add_body(doc, description)
    add_body(doc, f"Priority: {priority}", bold_prefix="Priority:")
    add_heading(doc, f"{number}.2 Stimulus/Response Sequences", 3)
    add_numbered(doc, sequence)
    add_heading(doc, f"{number}.3 Functional Requirements", 3)
    add_requirement_table(doc, requirements)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
styles["Normal"].font.size = Pt(11)
for style_name in ["Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    styles[style_name].font.name = "Times New Roman"
    styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

styles["Heading 1"].font.size = Pt(16)
styles["Heading 1"].font.bold = True
styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
styles["Heading 2"].font.size = Pt(14)
styles["Heading 2"].font.bold = True
styles["Heading 2"].font.color.rgb = RGBColor(46, 116, 181)
styles["Heading 3"].font.size = Pt(12)
styles["Heading 3"].font.bold = True
styles["Heading 3"].font.color.rgb = RGBColor(31, 78, 121)

for current_section in doc.sections:
    add_page_number(current_section.footer.paragraphs[0])

# Cover page
cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover.paragraph_format.space_before = Pt(90)
run = cover.add_run("SOFTWARE REQUIREMENTS\nSPECIFICATION")
run.bold = True
run.font.name = "Times New Roman"
run.font.size = Pt(24)
run.font.color.rgb = RGBColor(31, 78, 121)

project = doc.add_paragraph()
project.alignment = WD_ALIGN_PARAGRAPH.CENTER
project.paragraph_format.space_before = Pt(28)
run = project.add_run(f"for\n{PROJECT_NAME}")
run.bold = True
run.font.size = Pt(22)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_before = Pt(12)
run = subtitle.add_run("Blockchain-Based Bug Reporting and Tracking System")
run.italic = True
run.font.size = Pt(14)

details = doc.add_paragraph()
details.alignment = WD_ALIGN_PARAGRAPH.CENTER
details.paragraph_format.space_before = Pt(55)
details.add_run(
    f"Version {VERSION}\n\n"
    f"Prepared by\n{AUTHOR}\n\n"
    f"{ORGANIZATION}\n\n"
    f"{DATE}"
)

doc.add_page_break()

# Revision history
add_heading(doc, "Revision History", 1)
revision = doc.add_table(rows=2, cols=4)
revision.style = "Table Grid"
revision.alignment = WD_TABLE_ALIGNMENT.CENTER
for index, text in enumerate(["Name", "Date", "Reason for Changes", "Version"]):
    revision.rows[0].cells[index].text = text
    set_cell_shading(revision.rows[0].cells[index], "D9EAF7")
    for run in revision.rows[0].cells[index].paragraphs[0].runs:
        run.bold = True
for index, text in enumerate([AUTHOR, DATE, "Initial complete SRS for the BlockBug platform", VERSION]):
    revision.rows[1].cells[index].text = text

add_heading(doc, "Table of Contents", 1)
toc_paragraph = doc.add_paragraph()
add_toc(toc_paragraph)
doc.add_page_break()

# 1 Introduction
add_heading(doc, "1. Introduction", 1)
add_heading(doc, "1.1 Purpose", 2)
add_body(
    doc,
    "This Software Requirements Specification describes the requirements of BlockBug version 1.0. "
    "BlockBug is a web-based bug reporting and tracking platform that supports the complete bug life cycle, "
    "starting from bug reporting and ending at verification and closure. The system also uses blockchain "
    "technology to keep important bug events transparent and immutable. This document defines the functional "
    "requirements, interfaces, user roles, quality requirements, business rules, and deployment environment of the platform."
)
add_body(
    doc,
    "The SRS covers the frontend application, PHP backend API, PostgreSQL database, role-based workflows, "
    "project and sprint management, notifications, attachments, reports, and the blockchain audit component. "
    "It is intended to guide development, testing, demonstration, and future improvement of the system."
)

add_heading(doc, "1.2 Document Conventions", 2)
add_bullets(doc, [
    "The word shall is used for a mandatory requirement.",
    "The word should is used for a recommended behavior.",
    "Functional requirements are given unique identifiers such as AUTH-01 and BUG-01.",
    "Requirement priorities are written as High, Medium, or Low.",
    "Technical terms and abbreviations are defined in Appendix A.",
    "The document follows the main structure of the IEEE Software Requirements Specification template.",
])

add_heading(doc, "1.3 Intended Audience and Reading Suggestions", 2)
add_body(
    doc,
    "This document is mainly written for the project supervisor, students, developers, testers, project managers, "
    "and future maintainers of BlockBug. A reader who wants a general understanding should first read Sections 1 and 2. "
    "Developers should focus on Sections 3, 4, and 6. Testers can use the functional requirements in Section 4 as a base "
    "for test cases. The project supervisor and evaluators may read the product scope, gap analysis, architecture, "
    "nonfunctional requirements, and analysis models to understand the academic and practical value of the project."
)

add_heading(doc, "1.4 Product Scope", 2)
add_body(
    doc,
    "BlockBug is designed to help software teams report, assign, monitor, resolve, and verify bugs in one organized platform. "
    "Many development teams use issue trackers, spreadsheets, chat messages, or separate tools for bug handling. In such setups, "
    "the history of a bug may become difficult to verify because users with sufficient access can edit records, activities can be "
    "spread across different tools, and the final status may not clearly show who performed each important action."
)
add_body(
    doc,
    "The purpose of BlockBug is to combine normal bug management with a blockchain-supported audit trail. The regular application "
    "data is stored in PostgreSQL for fast searching and daily use, while important life-cycle events can be recorded on the Ethereum "
    "Sepolia test network through Alchemy. This hybrid approach keeps the system practical while providing stronger transparency "
    "for selected events such as bug creation, status change, resolution verification, and rejection of a fix."
)
add_bullets(doc, [
    "Provide one workspace for organizations, projects, team members, sprints, bugs, comments, and reports.",
    "Give separate responsibilities to administrators, managers, developers, and testers.",
    "Improve accountability by recording who reported, assigned, fixed, and verified a bug.",
    "Provide blockchain proof for important bug events without storing all application data on-chain.",
    "Offer dashboards and reports for understanding project quality and team activity.",
    "Support deployment as an online application that can be accessed using a modern web browser.",
])

add_heading(doc, "1.5 References", 2)
add_two_column_table(doc, [
    ("IEEE SRS Template", "Software Requirements Specification template by Karl E. Wiegers, 1999."),
    ("Next.js Documentation", "Documentation for the Next.js web application framework."),
    ("React Documentation", "Documentation for the React user interface library."),
    ("PHP Documentation", "Official PHP language and PDO documentation."),
    ("PostgreSQL Documentation", "Official PostgreSQL database documentation."),
    ("Hardhat Documentation", "Documentation for smart-contract development and testing."),
    ("Solidity Documentation", "Documentation for the Solidity smart-contract language."),
    ("Alchemy Documentation", "Documentation for connecting applications to Ethereum Sepolia."),
    ("Ethereum Sepolia", "Ethereum public test network used for blockchain deployment and testing."),
    ("Project Repository", "BlockBug source-code repository maintained on GitHub."),
], headers=("Reference", "Description"))

# 2 Overall Description
add_heading(doc, "2. Overall Description", 1)
add_heading(doc, "2.1 Product Perspective", 2)
add_body(
    doc,
    "BlockBug is a new, self-contained web platform made for software bug reporting and tracking. It is not only a simple form "
    "for entering bugs. It provides an organization-based workspace where different users take part in the bug life cycle according "
    "to their roles. The platform follows a layered architecture consisting of a browser-based frontend, a REST-style backend API, "
    "a PostgreSQL database, and a blockchain audit service."
)

add_heading(doc, "2.1.1 Problem Background", 3)
add_body(
    doc,
    "Bug reporting is an important activity in software development because defects must be communicated clearly before they can be fixed. "
    "A bug report usually contains a title, description, reproduction steps, expected result, actual result, environment, severity, and evidence. "
    "After reporting, the bug may pass through managers, developers, and testers. If this flow is not properly controlled, information can be lost, "
    "responsibility can become unclear, and the team may not know whether a bug was actually verified."
)
add_body(
    doc,
    "Conventional bug trackers are effective for daily project management, but their audit history normally remains inside one centralized database. "
    "The organization must trust the database administrators and the application permissions to protect that history. BlockBug keeps the advantages "
    "of a normal database while adding a separate blockchain proof layer for important actions."
)

add_heading(doc, "2.1.2 Existing System Study", 3)
add_body(
    doc,
    "Existing systems such as Jira, Bugzilla, GitHub Issues, Trello-based workflows, and spreadsheet-based reporting support useful issue-management "
    "functions. However, they are mainly centralized systems. Their focus is task organization, while independent proof of important bug events is "
    "usually not the main purpose. Smaller teams may also use several disconnected tools, for example a spreadsheet for bugs, chat for discussion, "
    "and another tool for sprint planning."
)
add_two_column_table(doc, [
    ("Centralized record history", "The history is normally stored and controlled by one application database."),
    ("Scattered communication", "Bug details, comments, screenshots, and decisions may be shared through separate channels."),
    ("Limited independent verification", "A third party may not have a simple method to verify that a recorded event existed at a certain stage."),
    ("General-purpose workflows", "Some systems require heavy configuration before they match a clear manager-developer-tester workflow."),
    ("Weak connection between verification and proof", "Tester approval is recorded as an application action but not always supported by immutable evidence."),
], headers=("Common Situation", "Explanation"))

add_heading(doc, "2.1.3 Identified Gap", 3)
add_body(
    doc,
    "The main gap identified for this project is the lack of a simple bug-reporting platform that combines role-based software quality workflow with "
    "blockchain-backed proof. Existing bug trackers mainly solve collaboration and task-management problems. Blockchain applications mainly focus on "
    "finance, ownership, or supply-chain records. BlockBug connects these areas by applying blockchain to selected software defect events."
)
add_bullets(doc, [
    "A transparent record is needed for important changes in the bug life cycle.",
    "Tester verification should become a clearly visible and provable stage.",
    "The team should use one platform for bug reports, projects, sprints, assignments, comments, and reports.",
    "Blockchain use should not make normal application operations slow or expensive.",
    "Sensitive and detailed bug information should remain off-chain, while compact event proof is stored on-chain.",
])

add_heading(doc, "2.1.4 Proposed Solution and Improvement", 3)
add_body(
    doc,
    "BlockBug uses a hybrid design. The frontend and PostgreSQL database handle normal work such as login, forms, filtering, comments, reports, and attachments. "
    "For selected events, the backend sends a structured audit request to a blockchain service. The smart contract stores the bug identifier, action, actor "
    "information, metadata, transaction time, and a blockchain bug-chain identifier. As a result, the platform remains easy to use like a normal web application "
    "but gains an additional proof source that is separate from the main database."
)
add_two_column_table(doc, [
    ("Transparency", "Authorized users can view the blockchain synchronization status and transaction reference of a bug event."),
    ("Immutability", "Once a transaction is confirmed on Sepolia, the recorded event cannot be silently changed inside the application database."),
    ("Accountability", "Important actions are connected with an actor email and role in the audit event."),
    ("Traceability", "Events can be followed using the bug ID, bug-chain ID, event ID, and transaction hash."),
    ("Practical performance", "Regular information remains in PostgreSQL, so searches and dashboard operations stay efficient."),
    ("Independent verification", "A transaction can be checked using the Ethereum Sepolia network and an explorer."),
], headers=("Improvement", "How BlockBug Provides It"))

add_heading(doc, "2.2 Product Functions", 2)
add_bullets(doc, [
    "Create and validate organization workspaces.",
    "Authenticate organization members and display role-based dashboards.",
    "Manage administrators, managers, developers, and testers.",
    "Create and manage software projects.",
    "Plan sprints and move unfinished bugs to a backlog or another sprint.",
    "Create bug reports with detailed reproduction information and optional attachments.",
    "Assign bugs to developers and verification testers.",
    "Move bugs through open, in-progress, resolved, and closed states.",
    "Allow testers to approve a fix or return it for more work.",
    "Add comments and replies to bug discussions.",
    "Generate notifications, dashboard summaries, and analytical reports.",
    "Maintain system settings, notification preferences, integrations, and API keys.",
    "Record important bug events on Ethereum Sepolia through Alchemy.",
    "Display blockchain transaction status and audit history for bugs.",
])

add_heading(doc, "2.3 User Classes and Characteristics", 2)
add_two_column_table(doc, [
    ("Administrator", "Has the highest application responsibility. Manages users, roles, system settings, API keys, integrations, and maintenance functions."),
    ("Manager", "Creates projects and sprints, reports bugs, assigns developers and testers, manages the backlog, and monitors reports."),
    ("Developer", "Views assigned bugs, comments on work, and moves an assigned bug from open to in-progress and then to resolved."),
    ("Tester", "Reports bugs, follows submitted reports, reviews resolved bugs, and either closes a verified fix or returns it to in-progress."),
    ("Project Supervisor or Evaluator", "May review the system workflow, reports, architecture, and blockchain audit evidence during project evaluation."),
])

add_heading(doc, "2.4 Operating Environment", 2)
add_two_column_table(doc, [
    ("Client device", "Desktop computer, laptop, tablet, or smartphone with a modern web browser."),
    ("Frontend", "Next.js and React application deployed on Vercel."),
    ("Backend", "PHP 8.2 REST-style API deployed as a Docker web service on Render."),
    ("Database", "PostgreSQL database connected to the backend."),
    ("Blockchain network", "Ethereum Sepolia test network accessed through an Alchemy RPC endpoint."),
    ("Smart contract tools", "Solidity, Hardhat, ethers.js, and Node.js audit service."),
    ("Network", "Internet connection using HTTPS for the deployed application and JSON-based API communication."),
])

add_heading(doc, "2.5 Design and Implementation Constraints", 2)
add_bullets(doc, [
    "The frontend shall use Next.js, React, and TypeScript.",
    "The backend shall use PHP and PDO for database communication.",
    "The primary database shall be PostgreSQL.",
    "The blockchain contract shall be written in Solidity and deployed to Ethereum Sepolia through Alchemy.",
    "The system shall use a browser-based responsive interface.",
    "Only selected audit information shall be stored on-chain; complete bug descriptions and attachments shall remain off-chain.",
    "The application shall follow the defined role-based bug workflow.",
    "The deployed frontend and backend shall communicate through HTTPS APIs.",
])

add_heading(doc, "2.6 User Documentation", 2)
add_bullets(doc, [
    "Software Requirements Specification document.",
    "Deployment notes for Vercel, Render, PostgreSQL, and environment variables.",
    "Blockchain workspace README with compile, deploy, and service commands.",
    "Simple in-application labels, validation messages, and role-based navigation.",
    "Project demonstration material and screenshots prepared for academic evaluation.",
])

add_heading(doc, "2.7 Assumptions and Dependencies", 2)
add_bullets(doc, [
    "Users have access to a modern browser and a stable internet connection.",
    "The Vercel frontend, Render backend, PostgreSQL database, Alchemy service, and Sepolia network are available.",
    "The organization administrator creates valid member accounts and assigns correct roles.",
    "The Sepolia wallet used by the blockchain service has enough test ETH to submit transactions.",
    "Environment variables are configured for the database, frontend URL, backend URL, blockchain RPC URL, contract address, and service key.",
    "The uploaded file type and size remain within the limits defined by the application.",
])

# 3 Interfaces
add_heading(doc, "3. External Interface Requirements", 1)
add_heading(doc, "3.1 User Interfaces", 2)
add_body(
    doc,
    "The user interface shall be a responsive web interface with consistent colors, cards, forms, tables, dialogs, buttons, and navigation elements. "
    "The application shall provide separate but related experiences for organization login, member login, signup, landing page, dashboard, bug reports, "
    "projects, reports, team management, and settings."
)
add_two_column_table(doc, [
    ("Landing Page", "Introduces the platform, main benefits, workflow, package information, and access links."),
    ("Organization Login", "Validates the organization ID and organization password before member login."),
    ("Member Login", "Accepts member email and password and opens the correct dashboard."),
    ("Dashboard", "Shows summary cards, recent activity, quick actions, charts, and role-specific information."),
    ("Bug Report Dialog", "Collects title, project, priority, severity, description, reproduction steps, results, environment, assignment details, and attachment."),
    ("Bug Detail View", "Shows complete bug data, comments, assignment, sprint, workflow actions, attachments, and blockchain audit events."),
    ("Project and Sprint Screens", "Provide project creation, sprint planning, backlog, completion, and carry-over actions."),
    ("Reports Screen", "Displays bug distribution, project statistics, and resolution information."),
    ("Team Screen", "Allows permitted users to view and manage team member accounts."),
    ("Settings Screen", "Provides profile, preferences, system settings, integrations, API keys, and maintenance options."),
])

add_heading(doc, "3.2 Hardware Interfaces", 2)
add_body(
    doc,
    "BlockBug does not require special hardware. The user interacts through a keyboard, mouse, touchpad, or touchscreen. "
    "The server environment requires normal cloud computing resources for the frontend, backend, database, and blockchain service. "
    "A device file picker is used when a user attaches an image, PDF, text file, CSV file, or ZIP file to a bug report."
)

add_heading(doc, "3.3 Software Interfaces", 2)
add_two_column_table(doc, [
    ("Next.js 16 and React 19", "Provide routing, rendering, component structure, and client-side interaction."),
    ("PHP 8.2 API", "Processes authentication, business rules, database actions, reports, file uploads, and blockchain requests."),
    ("PostgreSQL", "Stores organizations, users, projects, sprints, bugs, activities, comments, notifications, preferences, API keys, settings, and blockchain event references."),
    ("Hardhat and Solidity", "Compile, test, and deploy the BugAuditTrail smart contract."),
    ("ethers.js", "Connects the audit service to the Sepolia smart contract."),
    ("Alchemy Sepolia RPC", "Provides the blockchain connection used to submit and read Sepolia transactions."),
    ("Vercel", "Hosts the production frontend."),
    ("Render", "Hosts the PHP backend and connects it with the production database."),
])

add_heading(doc, "3.4 Communications Interfaces", 2)
add_bullets(doc, [
    "The browser and frontend shall use HTTPS.",
    "The frontend shall exchange JSON data with the backend API.",
    "Multipart form data shall be used for bug reports that include attachments.",
    "The backend shall communicate with PostgreSQL through the PostgreSQL protocol using PDO.",
    "The backend and blockchain audit service shall exchange JSON messages.",
    "The audit service shall communicate with Ethereum Sepolia through the Alchemy JSON-RPC endpoint.",
    "Blockchain transaction identifiers shall be stored in the application database for later viewing.",
])

# 4 System Features
add_heading(doc, "4. System Features", 1)

add_feature(doc, "4.1", "Organization and Member Authentication",
    "This feature controls entry into an organization workspace. A user first identifies the organization and then signs in using a member account. "
    "It provides a clear separation between different software organizations.",
    "High",
    [
        "The user enters the organization ID and organization password.",
        "The system validates the organization and displays the member login form.",
        "The user enters the member email and password.",
        "The system validates the account and opens the dashboard according to the member role.",
    ],
    [
        ("AUTH-01", "The system shall allow a new organization and its first administrator to be registered when public signup is enabled.", "High", "Functional test"),
        ("AUTH-02", "The system shall validate an organization ID and organization password.", "High", "API and UI test"),
        ("AUTH-03", "The system shall validate a member only within the selected organization.", "High", "API test"),
        ("AUTH-04", "The system shall reject an inactive organization or inactive member.", "High", "Negative test"),
        ("AUTH-05", "The system shall redirect an authenticated member to the dashboard.", "High", "UI test"),
        ("AUTH-06", "The system shall support logout and session timeout behavior.", "Medium", "UI test"),
    ])

add_feature(doc, "4.2", "Role-Based Workspace",
    "The system provides different menus and actions for administrators, managers, developers, and testers. Each role participates in the bug workflow in a different way.",
    "High",
    [
        "A member signs in successfully.",
        "The system reads the member role.",
        "The dashboard displays the pages, buttons, and workflow actions related to that role.",
    ],
    [
        ("ROLE-01", "The administrator shall be able to manage members, settings, integrations, API keys, and maintenance functions.", "High", "Role test"),
        ("ROLE-02", "The manager shall be able to manage projects, sprints, assignments, and reports.", "High", "Role test"),
        ("ROLE-03", "The developer shall see assigned bugs and permitted workflow actions.", "High", "Role test"),
        ("ROLE-04", "The tester shall be able to report bugs and verify resolved bugs assigned to the tester.", "High", "Role test"),
        ("ROLE-05", "The system shall hide actions that are not relevant to the current role.", "Medium", "UI test"),
    ])

add_feature(doc, "4.3", "Project and Sprint Management",
    "Projects organize related bugs, while sprints organize planned development work. Managers and administrators use these features to control current work and backlog movement.",
    "High",
    [
        "The manager creates a project.",
        "The manager creates a planned or active sprint with start and end dates.",
        "Bugs are assigned to the sprint or remain in the backlog.",
        "When completing a sprint, unfinished bugs are moved to the backlog or another valid sprint.",
    ],
    [
        ("PROJ-01", "The system shall allow permitted users to create a project with a unique project key inside the organization.", "High", "API and UI test"),
        ("PROJ-02", "The system shall list projects belonging to the current organization.", "High", "API test"),
        ("SPR-01", "The system shall create sprints with a name, goal, status, start date, and end date.", "High", "Functional test"),
        ("SPR-02", "The system shall prevent the sprint end date from being earlier than its start date.", "High", "Validation test"),
        ("SPR-03", "The system shall allow only one active sprint per project.", "Medium", "Business-rule test"),
        ("SPR-04", "The system shall record bug movement between sprints and the backlog.", "Medium", "Database test"),
        ("SPR-05", "The system shall require a carry-over decision for unfinished bugs when a sprint is completed.", "High", "Workflow test"),
    ])

add_feature(doc, "4.4", "Bug Reporting",
    "This feature allows managers and testers to create detailed bug reports. It captures enough information for a developer to understand and reproduce the problem.",
    "High",
    [
        "The user opens the new bug report dialog.",
        "The user enters the required and optional bug information.",
        "The user optionally selects an attachment.",
        "The system validates the report and stores it.",
        "The system creates activity, notification, and blockchain audit records for the new bug.",
    ],
    [
        ("BUG-01", "The system shall require a bug title, description, project, and reporter.", "High", "Validation test"),
        ("BUG-02", "The system shall support low, medium, high, and critical priorities.", "High", "Functional test"),
        ("BUG-03", "The system shall support minor, major, and critical severity values.", "High", "Functional test"),
        ("BUG-04", "The report shall support steps to reproduce, expected result, actual result, and environment.", "Medium", "UI and API test"),
        ("BUG-05", "The system shall allow supported attachments up to 10 MB.", "Medium", "Upload test"),
        ("BUG-06", "A manager-created bug shall include a developer and verification tester assignment.", "High", "Business-rule test"),
        ("BUG-07", "A tester-created bug shall use the reporting tester as the default verification tester.", "High", "Workflow test"),
        ("BUG-08", "The system shall show a clear validation or submission message when a report cannot be submitted.", "High", "UI test"),
    ])

add_feature(doc, "4.5", "Bug Assignment and Life-Cycle Workflow",
    "The bug life cycle makes responsibilities visible from reporting to closure. Developers perform development stages, while testers make the final verification decision.",
    "High",
    [
        "A manager assigns a bug to a developer and a tester.",
        "The assigned developer moves the bug from open to in-progress.",
        "After completing the fix, the developer moves the bug to resolved.",
        "The assigned tester checks the fix.",
        "The tester closes the verified bug or returns it to in-progress.",
    ],
    [
        ("FLOW-01", "The supported bug states shall be open, in-progress, resolved, and closed.", "High", "State test"),
        ("FLOW-02", "Only the assigned developer shall move an open bug to in-progress and an in-progress bug to resolved.", "High", "Authorization test"),
        ("FLOW-03", "Only the assigned or reporting tester shall verify a resolved bug.", "High", "Authorization test"),
        ("FLOW-04", "The tester shall be able to close a verified bug.", "High", "Workflow test"),
        ("FLOW-05", "The tester shall be able to return an unverified fix to in-progress.", "High", "Workflow test"),
        ("FLOW-06", "The system shall record status, assignment, tester, and sprint changes.", "High", "Database test"),
    ])

add_feature(doc, "4.6", "Comments, Replies, Attachments, and Notifications",
    "Collaboration features keep discussion and supporting material connected with the relevant bug. Notifications inform the team about important events.",
    "Medium",
    [
        "A user opens a bug detail view.",
        "The user adds a comment or replies to an existing comment.",
        "The system stores the message and creates a related activity and notification.",
        "Users can view attachments and notification history.",
    ],
    [
        ("COL-01", "The system shall allow members to add comments to a bug.", "Medium", "Functional test"),
        ("COL-02", "The system shall allow a comment to reply to another comment on the same bug.", "Medium", "Functional test"),
        ("COL-03", "The system shall display bug attachments with file name and size.", "Medium", "UI test"),
        ("NOT-01", "The system shall create notifications for important bug, project, sprint, user, and settings events.", "Medium", "Database test"),
        ("NOT-02", "The system shall support user notification preferences.", "Low", "Preferences test"),
        ("NOT-03", "The system shall provide notification targets that can open the related dashboard page.", "Medium", "Navigation test"),
    ])

add_feature(doc, "4.7", "Dashboard and Reporting",
    "The dashboard and reports provide a summarized view of software quality, team workload, project activity, and bug distribution.",
    "Medium",
    [
        "The user opens the dashboard or reports page.",
        "The system calculates statistics from the current organization data.",
        "The system displays totals, charts, recent activity, active users, and project information.",
    ],
    [
        ("REP-01", "The system shall display total, open, in-progress, resolved, closed, critical, and high-priority bug counts.", "Medium", "Data test"),
        ("REP-02", "The dashboard shall show weekly bug activity and comparison with the previous week.", "Medium", "Report test"),
        ("REP-03", "The system shall display active projects and average bugs per project.", "Medium", "Report test"),
        ("REP-04", "The system shall calculate average resolution information by project.", "Low", "Report test"),
        ("REP-05", "The tester dashboard shall show reported bugs and bugs waiting for verification.", "High", "Role UI test"),
    ])

add_feature(doc, "4.8", "Blockchain Audit Trail",
    "The blockchain audit feature provides an additional transparent and immutable record for important bug events. The smart contract is deployed on Ethereum Sepolia and accessed through Alchemy.",
    "High",
    [
        "An important bug event occurs in the application.",
        "The backend stores a pending blockchain event and sends audit data to the blockchain service.",
        "The service submits the transaction to the BugAuditTrail smart contract on Sepolia.",
        "The service returns the transaction hash, event ID, bug-chain ID, contract address, and block result.",
        "The backend marks the event as synchronized and displays the proof in the bug details.",
    ],
    [
        ("BC-01", "The system shall create a blockchain audit event for bug creation and important status changes.", "High", "Integration test"),
        ("BC-02", "The blockchain event shall include the bug ID, action, actor email, actor role, metadata, and blockchain timestamp.", "High", "Contract test"),
        ("BC-03", "The smart contract shall generate a stable bug-chain identifier for each bug.", "High", "Contract test"),
        ("BC-04", "The system shall store the Sepolia transaction hash, event ID, bug-chain ID, and contract address.", "High", "Database test"),
        ("BC-05", "The bug detail view shall show whether the blockchain event is pending, synchronized, or needs synchronization.", "Medium", "UI test"),
        ("BC-06", "The system shall support retrying an event that was not synchronized.", "Medium", "Integration test"),
        ("BC-07", "Detailed descriptions, credentials, passwords, and attachment files shall not be written to the blockchain.", "High", "Security review"),
        ("BC-08", "The blockchain service shall connect to Ethereum Sepolia using the configured Alchemy RPC endpoint.", "High", "Deployment test"),
    ])

add_feature(doc, "4.9", "Administration and Settings",
    "Administration features allow the organization to control users, system defaults, preferences, integrations, API keys, and maintenance operations.",
    "Medium",
    [
        "An administrator opens the settings or team page.",
        "The administrator changes a permitted configuration or member record.",
        "The system validates and stores the change.",
        "The system displays a success message and creates a related notification when required.",
    ],
    [
        ("ADM-01", "The administrator shall create, update, activate, deactivate, and remove member accounts.", "High", "Role test"),
        ("ADM-02", "The administrator shall configure default bug values, application name, timezone, date format, timeout, and signup setting.", "Medium", "Settings test"),
        ("ADM-03", "A user shall be able to update the user's own password after providing the current password.", "High", "Security test"),
        ("ADM-04", "The system shall support generating and revoking application API keys.", "Medium", "Functional test"),
        ("ADM-05", "The system shall display available integration options and their current status.", "Low", "UI test"),
        ("ADM-06", "The administrator shall be able to export organization maintenance data.", "Medium", "Export test"),
    ])

# 5 Nonfunctional
add_heading(doc, "5. Other Nonfunctional Requirements", 1)
add_heading(doc, "5.1 Performance Requirements", 2)
add_bullets(doc, [
    "Normal dashboard and list requests should respond within three seconds under expected academic demonstration load, excluding cloud cold-start time.",
    "The system shall use database indexes for organization, project, sprint, bug, and blockchain-event searches.",
    "The user interface shall display a loading or progress state during longer operations.",
    "Blockchain submission shall not prevent the user from viewing normal application data after the application event has been accepted.",
    "The system shall limit uploaded files to 10 MB to control server resource usage.",
    "Lists and reports should remain usable with at least several thousand bug records in one organization.",
])

add_heading(doc, "5.2 Safety Requirements", 2)
add_bullets(doc, [
    "Destructive maintenance actions shall be available only through clearly labeled administrative controls.",
    "The system shall ask for confirmation before deleting important records where applicable.",
    "Sprint completion shall require a decision for unfinished work so that bugs are not lost.",
    "The system shall preserve database relationships through foreign keys and controlled delete behavior.",
    "Uploaded files shall be checked for size and supported type before storage.",
])

add_heading(doc, "5.3 Security Requirements", 2)
add_bullets(doc, [
    "Organization and member passwords shall be stored as secure password hashes.",
    "Every protected API request shall be connected with an authenticated member session.",
    "The backend shall determine the user's organization and role from the authenticated session.",
    "The system shall enforce organization-level data separation.",
    "Only approved frontend origins shall be allowed to make browser requests to the backend.",
    "Role permissions shall be checked by the backend before a protected action is performed.",
    "Private blockchain keys and Alchemy credentials shall be stored in environment variables and shall not be returned to the browser.",
    "Production communication shall use HTTPS.",
    "Sensitive values shall not be included in error responses, logs, or blockchain metadata.",
    "API keys shall be stored as hashes, and the plain key shall only be shown when it is created.",
])

add_heading(doc, "5.4 Software Quality Attributes", 2)
add_two_column_table(doc, [
    ("Usability", "Forms, labels, messages, and role-based navigation shall be clear enough for a student or software-team member to use without long training."),
    ("Reliability", "Related database operations shall be treated as one complete transaction where partial data could otherwise be produced."),
    ("Maintainability", "Frontend, API, database, and blockchain code shall remain separated into understandable modules and documented interfaces."),
    ("Portability", "The system shall run through standard web technologies and container-based backend deployment."),
    ("Scalability", "Organization-scoped queries and indexed tables shall support growth in users, projects, bugs, and events."),
    ("Interoperability", "The platform shall communicate through JSON, HTTP, PostgreSQL, Ethereum JSON-RPC, and standard file upload formats."),
    ("Testability", "Functional requirements shall have clear identifiers and observable expected results."),
    ("Availability", "The deployed application should be available through Vercel, Render, PostgreSQL hosting, Alchemy, and Ethereum Sepolia."),
    ("Transparency", "Important bug events shall provide visible blockchain synchronization details."),
])

add_heading(doc, "5.5 Business Rules", 2)
add_bullets(doc, [
    "Each user belongs to an organization and has one application role.",
    "An organization project key must be unique inside that organization.",
    "Managers and testers may report bugs.",
    "A manager must select a developer and a tester when creating a bug.",
    "Only managers and administrators may create or complete sprints.",
    "Only one sprint may be active for a project at a time.",
    "A developer may move only an assigned bug through development states.",
    "A tester may verify only a resolved bug assigned to or reported by that tester.",
    "Closing a resolved bug represents successful tester verification.",
    "Returning a resolved bug to in-progress represents rejected verification.",
    "Important bug events shall create an application audit record and a blockchain audit request.",
])

# 6 Other requirements
add_heading(doc, "6. Other Requirements", 1)
add_heading(doc, "6.1 Database Requirements", 2)
add_body(
    doc,
    "The PostgreSQL database shall contain structured tables for organizations, users, projects, sprints, sprint movement history, bugs, attachments, "
    "comments, activities, notifications, preferences, system settings, API keys, integrations, and blockchain events. Foreign keys shall maintain "
    "important relationships, while organization identifiers shall keep workspace data separated."
)

add_heading(doc, "6.2 Blockchain Deployment Requirements", 2)
add_body(
    doc,
    "The BugAuditTrail contract is deployed on the Ethereum Sepolia test network, and the application connects to Sepolia through Alchemy. "
    "The blockchain service requires an Alchemy RPC URL, a funded Sepolia test wallet, the deployed contract address, and the compiled contract ABI. "
    "Transaction hashes produced by the service are linked back to bug records in the PostgreSQL database."
)

add_heading(doc, "6.3 Deployment Requirements", 2)
add_bullets(doc, [
    "The frontend is deployed on Vercel and is accessible through the BlockBug production website.",
    "The PHP backend is deployed on Render as a Docker-based web service.",
    "The backend uses a PostgreSQL database through the configured database connection URL.",
    "The frontend uses the public Render backend URL through the NEXT_PUBLIC_API_URL environment variable.",
    "The backend uses the Vercel URL as its approved frontend origin.",
    "The blockchain component uses Alchemy Sepolia environment configuration.",
])

add_heading(doc, "6.4 Data Retention and Audit Requirements", 2)
add_bullets(doc, [
    "Bug and project data shall remain available until removed by an authorized operation.",
    "Blockchain transaction references shall remain connected to the related bug event.",
    "Attachments shall be stored using persistent file or object storage in the deployed environment.",
    "Maintenance export shall provide organization data in a structured format.",
    "The application shall keep activity and sprint movement records for traceability.",
])

add_heading(doc, "6.5 Future Enhancement Possibilities", 2)
add_bullets(doc, [
    "Email and messaging notifications.",
    "Direct synchronization with GitHub, GitLab, Jira, Slack, and Microsoft Teams.",
    "Advanced search, pagination, and customizable dashboards.",
    "Smart-contract role permissions and multi-signature audit administration.",
    "Mainnet or private-network deployment after complete security and cost evaluation.",
    "Automated test execution and CI/CD quality gates.",
    "Artificial intelligence support for duplicate bug detection and report summarization.",
])

# Appendices
add_heading(doc, "Appendix A: Glossary", 1)
add_two_column_table(doc, [
    ("API", "Application Programming Interface used for communication between the frontend and backend."),
    ("Bug", "A defect or unexpected behavior found in software."),
    ("Bug Chain ID", "A blockchain identifier that groups audit events belonging to the same bug."),
    ("Blockchain", "A distributed ledger in which confirmed records are difficult to alter."),
    ("DApp", "A decentralized application or an application that uses blockchain components."),
    ("Immutability", "The property that a confirmed record cannot be silently changed."),
    ("Organization", "A separate workspace containing its own users, projects, bugs, and settings."),
    ("RPC", "Remote Procedure Call interface used to communicate with an Ethereum node."),
    ("Sepolia", "An Ethereum public test network used for development and demonstration."),
    ("Alchemy", "A blockchain infrastructure service used to access Ethereum Sepolia."),
    ("Smart Contract", "Program deployed on a blockchain that executes predefined functions."),
    ("Sprint", "A time-bounded period in which selected software work is planned."),
    ("Transaction Hash", "Unique identifier of a blockchain transaction."),
    ("Verification Tester", "Tester responsible for confirming whether a resolved bug is actually fixed."),
])

add_heading(doc, "Appendix B: Analysis Models", 1)
add_heading(doc, "B.1 High-Level Architecture Model", 2)
architecture = doc.add_table(rows=7, cols=1)
architecture.style = "Table Grid"
architecture.alignment = WD_TABLE_ALIGNMENT.CENTER
architecture_text = [
    "Users: Administrator | Manager | Developer | Tester",
    "↓ HTTPS through Web Browser",
    "Next.js + React Frontend on Vercel",
    "↓ JSON / Multipart REST Requests",
    "PHP 8.2 Backend API on Render ↔ PostgreSQL Database",
    "↓ Selected Audit Events",
    "Node.js Audit Service → Alchemy RPC → BugAuditTrail Contract on Ethereum Sepolia",
]
for index, text in enumerate(architecture_text):
    architecture.rows[index].cells[0].text = text
    architecture.rows[index].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if index in [0, 2, 4, 6]:
        set_cell_shading(architecture.rows[index].cells[0], "D9EAF7")
        for run in architecture.rows[index].cells[0].paragraphs[0].runs:
            run.bold = True
doc.add_paragraph()

add_heading(doc, "B.2 Bug State-Transition Model", 2)
state_table = doc.add_table(rows=1, cols=4)
state_table.style = "Table Grid"
state_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, text in enumerate(["Current State", "Action", "Responsible Role", "Next State"]):
    state_table.rows[0].cells[i].text = text
    set_cell_shading(state_table.rows[0].cells[i], "D9EAF7")
    for run in state_table.rows[0].cells[i].paragraphs[0].runs:
        run.bold = True
for row in [
    ("Open", "Start work", "Assigned Developer", "In Progress"),
    ("In Progress", "Mark fix ready", "Assigned Developer", "Resolved"),
    ("Resolved", "Verify fix", "Assigned/Reporting Tester", "Closed"),
    ("Resolved", "Reject fix", "Assigned/Reporting Tester", "In Progress"),
]:
    cells = state_table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
doc.add_paragraph()

add_heading(doc, "B.3 Main Data Entities", 2)
add_two_column_table(doc, [
    ("Organization", "Owns users, projects, bugs, activities, notifications, and sprints."),
    ("User", "Represents a member with a role, status, profile, preferences, and API keys."),
    ("Project", "Groups related bugs and sprints."),
    ("Sprint", "Plans work for a project and records completion or carry-over."),
    ("Bug", "Stores the defect, workflow state, assignments, verification details, and audit summary."),
    ("Comment", "Stores discussion and reply relationships for a bug."),
    ("Attachment", "Stores supporting file metadata for a bug."),
    ("Activity", "Stores readable application history."),
    ("Notification", "Informs users about important events."),
    ("Blockchain Event", "Stores synchronization status and Sepolia transaction references."),
])

add_heading(doc, "B.4 Simplified Data Flow", 2)
add_numbered(doc, [
    "The user performs an action in the Vercel-hosted frontend.",
    "The frontend sends a validated request to the Render-hosted PHP API.",
    "The backend applies role and business rules.",
    "The backend reads or writes organization data in PostgreSQL.",
    "For an important bug event, the backend creates an audit request.",
    "The audit service submits the event to the Sepolia smart contract through Alchemy.",
    "The transaction result is stored in PostgreSQL and displayed in the bug detail view.",
])

add_heading(doc, "Appendix C: To Be Determined List", 1)
add_body(
    doc,
    "The core functional and deployment requirements are defined for version 1.0. The following project details may be finalized during future releases:"
)
add_numbered(doc, [
    "The final production subscription packages and billing implementation.",
    "The long-term blockchain network choice after academic testing on Sepolia.",
    "The production object-storage provider for attachments.",
    "The exact retention period for notifications and activity records.",
    "The final list of third-party integrations enabled for organizations.",
])

# Header after cover
for current_section in doc.sections:
    header = current_section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run(f"Software Requirements Specification for {PROJECT_NAME}")
    header_run.font.size = Pt(9)
    header_run.font.italic = True
    header_run.font.color.rgb = RGBColor(100, 100, 100)

# Document properties and update-fields flag
doc.core_properties.title = f"Software Requirements Specification for {PROJECT_NAME}"
doc.core_properties.subject = "Blockchain-Based Bug Reporting and Tracking System"
doc.core_properties.author = AUTHOR
doc.core_properties.keywords = "SRS, BlockBug, bug reporting, blockchain, Sepolia, Alchemy"
doc.core_properties.comments = "Prepared using the IEEE-style SRS structure."

settings = doc.settings._element
update_fields = OxmlElement("w:updateFields")
update_fields.set(qn("w:val"), "true")
settings.append(update_fields)

OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT_PATH)
print(OUTPUT_PATH)
